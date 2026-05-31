import streamlit as st
import anthropic
import base64
import fitz  # PyMuPDF

st.set_page_config(
    page_title="QuantAI — Quantificação e Orçamento com IA",
    page_icon="📐",
    layout="wide",
)

st.markdown("""
<style>
  [data-testid="stAppViewContainer"] { background: #ffffff; }
  [data-testid="stHeader"] { background: #ffffff; }
  [data-testid="stSidebar"] { background: #fff8f3; border-right: 2px solid #ff8c3a; }
  .main .block-container { padding-top: 2rem; max-width: 1100px; }
  h1, h2, h3, h4, h5 { color: #1a1a1a !important; }
  p, li, label { color: #1a1a1a !important; }
  [data-testid="stMarkdownContainer"] h1 { color: #ff6b00 !important; }
  [data-testid="stMarkdownContainer"] h2 { color: #ff8c3a !important; }
  .stButton > button {
    background: #ff6b00 !important; color: #ffffff !important;
    font-weight: 700 !important; border: none !important;
    border-radius: 10px !important; padding: 0.75rem 2rem !important;
    font-size: 1rem !important; width: 100%;
  }
  .stButton > button:hover { background: #ff8c3a !important; }
  [data-testid="stFileUploader"] { border: 2px dashed #ff8c3a !important; border-radius: 10px !important; background: #fff8f3 !important; }
  [data-testid="stTextArea"] textarea { border: 1.5px solid #ffcca0 !important; border-radius: 8px !important; background: #fff8f3 !important; color: #1a1a1a !important; }
  [data-testid="stSelectbox"] > div { border: 1.5px solid #ffcca0 !important; border-radius: 8px !important; background: #fff8f3 !important; }
  [data-testid="stPasswordInput"] input { border: 1.5px solid #ffcca0 !important; background: #fff8f3 !important; color: #1a1a1a !important; }
  table { border-collapse: collapse; width: 100%; }
  thead tr { background: #ff6b00 !important; color: white !important; }
  thead th { color: white !important; padding: 10px 14px !important; font-weight: 600 !important; }
  tbody tr:nth-child(even) { background: #fff8f3; }
  tbody td { padding: 8px 14px !important; color: #1a1a1a !important; border-bottom: 1px solid #ffe0c8 !important; }
  footer { visibility: hidden; }
</style>
""", unsafe_allow_html=True)

# ── Sidebar ───────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown("## ⚡ QuantAI")
    st.markdown("Quantificação e orçamento de projetos com Inteligência Artificial.")
    st.divider()
    api_key = st.text_input("🔑 API Key (Anthropic)", type="password", placeholder="sk-ant-...", help="Obtenha em console.anthropic.com")
    st.divider()
    st.markdown("**Configurações**")
    currency = st.selectbox("Moeda", ["BRL — Real", "USD — Dólar", "EUR — Euro"])
    price_table = st.selectbox("Tabela de preços", ["SINAPI (Brasil)", "PINI", "Preços de mercado"])
    analysis_type = st.selectbox("Tipo de análise", ["Análise completa", "Só estrutura", "Só acabamentos", "Só instalações"])
    st.divider()
    st.caption("Zigurat — Trabalho M7_T2")

# ── Main ──────────────────────────────────────────────────────────────────────
st.markdown("# 📐 QuantAI")
st.markdown("##### Quantificação e Orçamento com Inteligência Artificial")
st.markdown("Envie a planta do seu projeto e receba automaticamente os quantitativos de materiais e o orçamento estimado.")
st.divider()

col1, col2 = st.columns([3, 2])
with col1:
    uploaded_file = st.file_uploader(
        "📁 Envie a planta do projeto",
        type=["pdf", "png", "jpg", "jpeg", "webp"],
        help="Planta baixa, corte, fachada ou qualquer imagem do projeto"
    )
with col2:
    context = st.text_area("💬 Contexto adicional (opcional)",
        placeholder="Ex: Apartamento de 80m² em SP. Reforma de banheiros e cozinha. Padrão médio.",
        height=120)

# ── Visualizador ──────────────────────────────────────────────────────────────
if uploaded_file is not None:
    file_bytes = uploaded_file.read()
    is_pdf = uploaded_file.type == "application/pdf"

    with col1:
        if is_pdf:
            try:
                doc = fitz.open(stream=file_bytes, filetype="pdf")
                page = doc[0]
                mat = fitz.Matrix(1.5, 1.5)
                pix = page.get_pixmap(matrix=mat)
                img_bytes = pix.tobytes("png")
                st.image(img_bytes, caption=f"Página 1 de {len(doc)}", use_container_width=True)
            except Exception as e:
                st.info("📄 PDF carregado com sucesso.")
        else:
            st.image(file_bytes, use_container_width=True)

    uploaded_file.seek(0)

st.markdown("")
analyze = st.button("⚡ Analisar com IA", disabled=not (uploaded_file and api_key))

if uploaded_file and not api_key:
    st.warning("Insira sua API Key da Anthropic na barra lateral para continuar.")

# ── Analysis ──────────────────────────────────────────────────────────────────
if analyze and uploaded_file and api_key:
    currency_code = currency.split(" — ")[0]
    symbols = {"BRL": "R$", "USD": "US$", "EUR": "€"}
    sym = symbols.get(currency_code, "R$")

    prompt = f"""Você é um engenheiro de custos especialista em orçamentos de construção civil no Brasil.

Analise o projeto enviado e gere um orçamento detalhado.

{f'Contexto: {context}' if context else ''}
Moeda: {currency_code} ({sym})
Tabela de referência: {price_table}
Escopo: {analysis_type}

Responda em markdown com estas seções exatamente nesta ordem:

## 📋 Resumo do Projeto
Uma descrição curta do projeto identificado, área estimada, número de ambientes e padrão de acabamento.

## 🏠 Ambientes Identificados
Uma tabela markdown com colunas: Ambiente | Área (m²) | Perímetro (m) | Pé-direito (m)

## 💰 Quantitativos e Orçamento
Uma tabela markdown com colunas: Categoria | Item | Especificação | Qtd | Un | Custo Unit. ({sym}) | Total ({sym})
Ao final da tabela, adicione uma linha de TOTAL GERAL em negrito.

## 💡 Observações
Premissas adotadas e recomendações importantes.

Use valores realistas baseados em {price_table}."""

    file_bytes = uploaded_file.read()
    file_b64 = base64.standard_b64encode(file_bytes).decode("utf-8")
    is_pdf = uploaded_file.type == "application/pdf"

    with st.spinner("🔍 Analisando o projeto com IA... isso pode levar até 30 segundos."):
        try:
            client = anthropic.Anthropic(api_key=api_key)
            if is_pdf:
                content_blocks = [
                    {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": file_b64}},
                    {"type": "text", "text": prompt}
                ]
                response = client.beta.messages.create(
                    model="claude-opus-4-5", max_tokens=8000,
                    messages=[{"role": "user", "content": content_blocks}],
                    betas=["pdfs-2024-09-25"]
                )
            else:
                media_type = uploaded_file.type or "image/jpeg"
                content_blocks = [
                    {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": file_b64}},
                    {"type": "text", "text": prompt}
                ]
                response = client.messages.create(
                    model="claude-opus-4-5", max_tokens=8000,
                    messages=[{"role": "user", "content": content_blocks}]
                )
            st.session_state["result"] = response.content[0].text
        except anthropic.AuthenticationError:
            st.error("❌ API Key inválida.")
            st.stop()
        except anthropic.RateLimitError:
            st.error("❌ Limite de requisições atingido. Aguarde e tente novamente.")
            st.stop()
        except Exception as e:
            st.error(f"❌ Erro: {str(e)}")
            st.stop()

# ── Results ───────────────────────────────────────────────────────────────────
if "result" in st.session_state:
    st.divider()
    st.success("✅ Análise concluída!")
    st.markdown(st.session_state["result"])
    st.divider()
    import io, re
    md = st.session_state["result"]

    # Parse markdown tables into Excel sheets
    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Orçamento"

        orange_fill = PatternFill("solid", fgColor="FF6B00")
        light_fill = PatternFill("solid", fgColor="FFF8F3")
        white_fill = PatternFill("solid", fgColor="FFFFFF")
        bold_font = Font(bold=True, color="FFFFFF")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        title_font = Font(bold=True, size=13, color="FF6B00")
        thin = Side(style="thin", color="FFE0C8")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        row = 1
        for line in md.split("\n"):
            line = line.strip()
            if line.startswith("## "):
                row += 1
                cell = ws.cell(row=row, column=1, value=line.replace("## ", "").replace("#","").strip())
                cell.font = title_font
                row += 1
            elif line.startswith("|") and not line.startswith("|---"):
                cells = [c.strip() for c in line.split("|")[1:-1]]
                for col, val in enumerate(cells, 1):
                    c = ws.cell(row=row, column=col, value=val)
                    c.border = border
                    c.alignment = Alignment(wrap_text=True, vertical="center")
                    if row > 1:
                        prev = ws.cell(row=row-1, column=1).value
                        # Check if previous row was a header (orange fill)
                        if ws.cell(row=row-1, column=col).fill.fgColor.rgb == "FFFF6B00":
                            c.fill = light_fill if (row % 2 == 0) else white_fill
                        else:
                            # header row
                            c.fill = orange_fill
                            c.font = header_font
                    else:
                        c.fill = orange_fill
                        c.font = header_font
                row += 1
            elif line.startswith("|---"):
                continue

        # Auto width
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        st.download_button(
            label="⬇ Baixar Orçamento (.xlsx)",
            data=buf,
            file_name="orcamento_quantai.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    except Exception as e:
        st.download_button(
            label="⬇ Baixar resultado (.md)",
            data=md,
            file_name="orcamento_quantai.md",
            mime="text/markdown"
        )
